from __future__ import annotations

import io

from fastapi import HTTPException
from openpyxl import load_workbook

import app.inventory_view_export as export


def _kwargs(**overrides):
    values = {
        "preset": "cms-mapping-review",
        "fields": "local_item_name,cms_code",
        "q": "tablet",
        "mapping_status": "REVIEW_REQUIRED",
        "source_classification": None,
        "review_reason": "continuity",
        "sort_field": "local_item_name",
        "sort_dir": "desc",
    }
    values.update(overrides)
    return values


def main() -> None:
    original_export_rows = export._export_rows
    try:
        captured: dict[str, object] = {}

        def fake_rows(**kwargs):
            captured.update(kwargs)
            return [
                {"local_item_name": "=SUM(A1:A2)", "cms_code": "@CMS-1"},
                {"local_item_name": "Normal, Item", "cms_code": "CMS-2"},
            ]

        export._export_rows = fake_rows

        xlsx_response = export.inventory_view_export_xlsx(**_kwargs())
        assert xlsx_response.status_code == 200
        assert xlsx_response.headers["content-disposition"] == 'attachment; filename="msa-cms-mapping-review.xlsx"'
        assert xlsx_response.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        assert xlsx_response.headers["cache-control"] == "no-store"
        assert xlsx_response.headers["x-msa-export-read-only"] == "true"
        assert xlsx_response.headers["x-msa-database-canonical"] == "false"
        assert xlsx_response.headers["x-msa-migration-baseline-accepted"] == "false"
        workbook = load_workbook(io.BytesIO(xlsx_response.body), data_only=False)
        worksheet = workbook["CMS Mapping Review"]
        assert worksheet.freeze_panes == "A2"
        assert worksheet["A1"].value == "Local Item"
        assert worksheet["B1"].value == "CMS Code"
        assert worksheet["A2"].value == "'=SUM(A1:A2)"
        assert worksheet["B2"].value == "'@CMS-1"
        assert worksheet["A1"].font.bold is True
        assert worksheet["A1"].alignment.wrap_text is True
        assert worksheet["A2"].border.left.style == "thin"
        assert len(worksheet.tables) == 1

        csv_response = export.inventory_view_export_csv(**_kwargs())
        text = csv_response.body.decode("utf-8")
        assert text.startswith("\ufeffLocal Item,CMS Code\n")
        assert "'=SUM(A1:A2),'@CMS-1" in text
        assert '"Normal, Item",CMS-2' in text

        assert captured["provider"] == "cms_mapping_review"
        assert captured["q"] == "tablet"
        assert captured["mapping_status"] == "REVIEW_REQUIRED"
        assert captured["review_reason"] == "continuity"
        assert captured["sort_field"] == "local_item_name"
        assert captured["sort_dir"] == "desc"

        try:
            export.inventory_view_export_xlsx(**_kwargs(preset="main-stock", fields="local_item_name,raw_sql_expression", q=None, mapping_status=None, review_reason=None, sort_field=None, sort_dir=None))
        except HTTPException as exc:
            assert exc.status_code == 400
        else:
            raise AssertionError("unknown export field must be rejected")

        try:
            export.inventory_view_export_xlsx(**_kwargs(fields="local_item_name", q=None, mapping_status=None, review_reason=None, sort_field="raw_sql_expression"))
        except HTTPException as exc:
            assert exc.status_code == 400
        else:
            raise AssertionError("unknown export sort field must be rejected")

        export._export_rows = lambda **_: [{} for _ in range(export.MAX_EXPORT_ROWS + 1)]
        try:
            export.inventory_view_export_xlsx(**_kwargs(preset="main-stock", fields="local_item_name", q=None, mapping_status=None, review_reason=None, sort_field=None, sort_dir=None))
        except HTTPException as exc:
            assert exc.status_code == 422
            assert str(export.MAX_EXPORT_ROWS) in exc.detail
        else:
            raise AssertionError("Excel export must reject rows beyond the hard cap")
    finally:
        export._export_rows = original_export_rows

    paths = {getattr(route, "path", None) for route in export.router.routes}
    assert "/dashboard/api/inventory-view/export.xlsx" in paths
    assert "/dashboard/api/inventory-view/export.csv" in paths
    assert export.MAX_EXPORT_ROWS == 5000
    print(
        "inventory_excel_export=pass reusable_renderer=pass registered_fields=pass validated_sort=pass "
        "formatted_xlsx=pass formula_literal=pass row_cap=5000 no_store=pass read_only=pass csv_compat=pass"
    )


if __name__ == "__main__":
    main()
