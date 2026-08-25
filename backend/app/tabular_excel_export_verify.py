from __future__ import annotations

import io
from datetime import date
from decimal import Decimal

from openpyxl import load_workbook

from app.tabular_excel_export import ExcelColumn, ExcelWorkbookSpec, build_excel_workbook


def main() -> None:
    payload = build_excel_workbook(
        ExcelWorkbookSpec(
            sheet_name="Reusable Export",
            table_name="Reusable Export Table",
            columns=(
                ExcelColumn("item", "Items", "string", 26),
                ExcelColumn("qty", "Current Qty", "decimal", 14, "0"),
                ExcelColumn("price", "Price", "decimal", 14, "0.00"),
                ExcelColumn("expiry", "Expiry Date", "date", 14, "mmm-yy"),
                ExcelColumn("note", "Review Reason", "string", 32),
            ),
            rows=(
                {
                    "item": "=SUM(A1:A2)",
                    "qty": Decimal("12"),
                    "price": Decimal("12.500"),
                    "expiry": date(2027, 1, 31),
                    "note": "Long wrapped text that should remain readable in the generated workbook without changing the underlying data selection contract.",
                },
                {"item": "Normal Item", "qty": Decimal("3"), "price": Decimal("7.25"), "expiry": None, "note": "@literal"},
            ),
        )
    )
    assert payload[:2] == b"PK"

    workbook = load_workbook(io.BytesIO(payload), data_only=False)
    worksheet = workbook["Reusable Export"]
    assert worksheet.freeze_panes == "A2"
    assert worksheet.sheet_view.showGridLines is True
    assert worksheet["A1"].value == "Items"
    assert worksheet["A1"].fill.fgColor.rgb in {"001F4E78", "1F4E78"}
    assert worksheet["A1"].font.bold is True
    assert worksheet["A1"].font.color.type == "rgb"
    assert worksheet["A1"].font.color.rgb in {"00FFFFFF", "FFFFFFFF"}
    assert worksheet["A2"].value == "'=SUM(A1:A2)"
    assert worksheet["E3"].value == "'@literal"
    assert worksheet["B2"].value == 12
    assert worksheet["B2"].number_format == "0"
    assert worksheet["C2"].value == 12.5
    assert worksheet["C2"].number_format == "0.00"
    assert worksheet["D2"].number_format == "mmm-yy"
    assert worksheet["A2"].alignment.wrap_text is True
    assert worksheet["A2"].border.left.style == "thin"
    assert worksheet.column_dimensions["A"].width == 26
    assert worksheet.row_dimensions[2].height >= 18
    assert len(worksheet.tables) == 1
    table = next(iter(worksheet.tables.values()))
    assert table.ref == "A1:E3"
    assert table.autoFilter is not None
    print(
        "reusable_excel_export=pass xlsx=pass header_style=pass freeze_panes=pass "
        "wrap=pass sizing=pass borders=pass table_filter=pass typed_cells=pass "
        "format_overrides=pass formula_literal=pass"
    )


if __name__ == "__main__":
    main()
